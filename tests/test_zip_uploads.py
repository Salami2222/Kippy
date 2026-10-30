"""Archive security, partial imports, ordinary document reuse and course isolation."""
import io
import stat
import struct
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from kippy import documents
from kippy.storage import Store
from kippy.web import App
import test_kippy as fixtures


def archive_bytes(entries,compression=zipfile.ZIP_DEFLATED):
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',compression=compression) as archive:
        for name,content in entries:archive.writestr(name,content)
    return buffer.getvalue()


class ZipUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.store=Store(self.temp.name)
        self.course=self.store.create_course({'name':'Math'})
    def tearDown(self):self.temp.cleanup()
    def upload(self,content,limit=25*1024*1024):
        return documents.upload(self.store,'lectures.zip',content,limit,self.course['id'])
    def assert_rejected(self,content,pattern):
        with self.assertRaisesRegex(ValueError,pattern):self.upload(content)
        self.assertEqual(self.store.materials(),[]);self.assertEqual(list(self.store.uploads.iterdir()),[])
    def test_multiple_pdfs_are_individual_citable_materials(self):
        data=archive_bytes([('Week 1/Lecture.pdf',fixtures.pdf_bytes(['Wronskian first page','Second page'])),('Week 2/Lecture.pdf',fixtures.pdf_bytes(['Variation of parameters']))])
        result=self.upload(data);self.assertEqual(result['summary'],{'imported':2,'failed':0,'skipped':0})
        docs=result['materials'];self.assertEqual([d['name'] for d in docs],['Lecture.pdf','Lecture.pdf']);self.assertNotEqual(docs[0]['id'],docs[1]['id'])
        self.assertEqual(len(self.store.materials()),2);self.assertEqual({c['location'] for c in self.store.course_chunks(self.course['id'])},{'page 1','page 2'})
        self.assertEqual([m['type'] for m in docs],['pdf','pdf']);self.assertEqual(len(list(self.store.uploads.iterdir())),2)
    def test_mixed_types_nested_directories_metadata_unsupported_and_partial_failure(self):
        from pptx import Presentation
        from docx import Document
        slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='Eigenvalues';pptx=io.BytesIO();slides.save(pptx)
        word=Document();word.add_paragraph('Integrating factors');docx=io.BytesIO();word.save(docx)
        data=archive_bytes([('Notes/',b''),('Notes/a.txt',b'Wronskian'),('Notes/slides.pptx',pptx.getvalue()),('Notes/word.docx',docx.getvalue()),('bad.pdf',b'not a PDF'),('nested.zip',b'not extracted'),('image.png',b'unsupported'),('__MACOSX/._a.txt',b'junk'),('.DS_Store',b'junk')])
        r=self.upload(data);self.assertEqual(r['summary'],{'imported':3,'failed':1,'skipped':2});self.assertEqual(len(r['materials']),4)
        self.assertIn('not a PDF',r['failed'][0]['error']);self.assertEqual(self.store.material(r['failed'][0]['material_id'])['status'],'failed')
        self.assertEqual({c['location'] for c in self.store.course_chunks(self.course['id'])},{'lines 1–1','slide 1','section 1'})
    def test_duplicates_never_overwrite_existing_and_all_chats_share_only_ready_imports(self):
        old=documents.upload(self.store,'same.txt',b'Existing',1000,self.course['id'])
        r=self.upload(archive_bytes([('one/same.txt',b'NEW_SOURCE_ONE'),('two/same.txt',b'NEW_SOURCE_TWO'),('broken.pdf',b'broken')]))
        split=self.store.split_course(self.course['id'],'a'*32);other=self.store.create_course({'name':'Chemistry'})
        ready={old['id']}|{d['material_id'] for d in r['imported']}
        restored=Store(self.temp.name)
        for key in ['left_id','right_id']:self.assertEqual(set(restored.course_conversation(self.course['id'],split['workspace'][key])['material_ids']),ready)
        self.assertEqual(restored.course_chunks(other['id']),[])
        client=fixtures.FakeClient();app=App(self.temp.name,client)
        for key in ['left_id','right_id']:app.chat({'course_id':self.course['id'],'conversation_id':split['workspace'][key],'request_id':'b'*32,'message':'Explain NEW_SOURCE_ONE and NEW_SOURCE_TWO'})
        for call in client.calls:self.assertIn('NEW_SOURCE_ONE',str(call));self.assertIn('NEW_SOURCE_TWO',str(call))
        removed=r['imported'][0]['material_id'];restored.delete_course_material(self.course['id'],removed)
        self.assertFalse((restored.uploads/removed).exists());self.assertEqual((restored.uploads/old['id']).read_bytes(),b'Existing')
        self.assertEqual(len(restored.course(self.course['id'])['materials']),3)
    def test_unsafe_paths_symlinks_and_special_entries_reject_everything_before_import(self):
        for path in ['../outside.txt','/absolute.txt','C:/drive.txt',r'..\outside.txt','folder/../../outside.txt','file.txt:stream','bad\x00.txt']:
            with self.subTest(path=path):
                info=zipfile.ZipInfo(path)
                if '\x00' in path:info.orig_filename=path;info.filename=path  # zipfile writer preserves the malformed name.
                self.assert_rejected(archive_bytes([('valid.txt',b'Keep nothing'),(info,b'unsafe')]),'unsafe')
        for mode in [stat.S_IFLNK,stat.S_IFIFO,stat.S_IFCHR]:
            info=zipfile.ZipInfo('special.txt');info.create_system=3;info.external_attr=(mode|0o777)<<16
            self.assert_rejected(archive_bytes([('valid.txt',b'Keep nothing'),(info,b'target')]),'unsafe')
    def test_encrypted_archive_is_rejected_before_importing_other_entries(self):
        data=bytearray(archive_bytes([('valid.txt',b'Good'),('locked.txt',b'Protected')]))
        local=data.find(b'PK\x03\x04',4);central=data.find(b'PK\x01\x02');central=data.find(b'PK\x01\x02',central+4)
        struct.pack_into('<H',data,local+6,struct.unpack_from('<H',data,local+6)[0]|1)
        struct.pack_into('<H',data,central+8,struct.unpack_from('<H',data,central+8)[0]|1)
        self.assert_rejected(bytes(data),'Encrypted')
    def test_corrupt_archives_and_zip64_or_multidisk_are_rejected(self):
        self.assert_rejected(b'not a ZIP','valid ZIP')
        good=archive_bytes([('a.txt',b'Alpha')]);self.assert_rejected(good[:-10],'valid ZIP')
        for offset,value,pattern in [(4,1,'multi-disk'),(10,65535,'ZIP64'),(16,0xffffffff,'ZIP64')]:
            data=bytearray(good);end=data.rfind(b'PK\x05\x06');struct.pack_into('<H' if offset<12 else '<L',data,end+offset,value)
            # ZIP64 sentinel entry counts must agree to reach the ZIP64-specific rejection.
            if offset==10:struct.pack_into('<H',data,end+8,value)
            self.assert_rejected(bytes(data),pattern)
    def test_file_count_total_size_ratio_and_directory_count_limits(self):
        self.assert_rejected(archive_bytes([(str(i)+'.txt',b'Notes') for i in range(51)]),'50 supported')
        with patch.object(documents,'MAX_ARCHIVE_BYTES',10):self.assert_rejected(archive_bytes([('a.txt',b'123456'),('b.txt',b'123456')]),'100 MiB')
        self.assert_rejected(archive_bytes([('bomb.txt',b'A'*(2*1024*1024))]),'compression ratio')
        with patch.object(documents,'MAX_ARCHIVE_ENTRIES',2):self.assert_rejected(archive_bytes([('a.exe',b'A'),('b.exe',b'B'),('c.exe',b'C')]),'directory')
    def test_individual_limit_partial_crc_failure_and_unsupported_only(self):
        # A high compressed-size limit can coexist with a smaller extracted-size cap in the importer.
        data=archive_bytes([('small.txt',b'A'),('large.txt',b'1234567890')])
        r=documents.import_archive(self.store,'test.zip',data,5,self.course['id']);self.assertEqual(r['summary']['imported'],1);self.assertEqual(r['summary']['failed'],1)
        corrupt=bytearray(archive_bytes([('good.txt',b'Good source'),('corrupt.txt',b'Wrong bytes')],zipfile.ZIP_STORED));at=corrupt.find(b'Wrong bytes');corrupt[at]=ord('X')
        r=self.upload(bytes(corrupt));self.assertEqual(r['summary']['imported'],1);self.assertEqual(r['summary']['failed'],1);self.assertIn('CRC',r['failed'][0]['error'])
        r=self.upload(archive_bytes([('nested.zip',b'ignored'),('picture.jpg',b'ignored')]))
        self.assertEqual(r['summary'],{'imported':0,'failed':0,'skipped':2})
    def test_compressed_upload_limit_and_ordinary_files_still_use_original_response(self):
        with self.assertRaisesRegex(ValueError,'upload size'):self.upload(archive_bytes([('a.txt',b'Notes')]),5)
        result=documents.upload(self.store,'plain.txt',b'Notes',100,self.course['id']);self.assertEqual(result['status'],'ready');self.assertNotIn('kind',result)


class ZipHTTPTests(unittest.TestCase):
    setUp=fixtures.HTTPTests.setUp
    tearDown=fixtures.HTTPTests.tearDown
    request=fixtures.HTTPTests.request
    def test_course_upload_route_returns_batch_summary_and_persisted_individual_files(self):
        _,course=self.request('/api/courses',{'name':'Math'});root='/api/courses/'+course['id']
        status,result=self.request(root+'/materials?filename=lectures.zip',archive_bytes([('a.txt',b'Calculus'),('b.txt',b'Equations')]),headers={'Content-Type':'application/octet-stream'})
        self.assertEqual(status,201);self.assertEqual(result['kind'],'archive');self.assertEqual(result['summary']['imported'],2)
        self.assertEqual(len(self.request(root)[1]['materials']),2)
        self.assertEqual(self.request(root+'/materials?filename=bad.zip',b'broken',headers={'Content-Type':'application/octet-stream'})[0],400)
        self.assertEqual(len(self.request(root)[1]['materials']),2)
